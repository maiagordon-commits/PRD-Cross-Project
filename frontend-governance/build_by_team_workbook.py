#!/usr/bin/env python3
"""Build Frontend Governance violations workbook with correct team name mapping."""

from __future__ import annotations

import csv
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = Path(
    "/home/ubuntu/.cursor/projects/workspace/uploads/"
    "Frontend_Governance__2026-07-26__-_Data_Rule_Violations_465c.csv"
)

# GitHub owner slug → Guesty team name (per product/eng naming)
# Explicit corrections from Maia (2026-07-29):
#   biz-enterprise → Biz
#   discovery → Odyssey
#   okta-communication → Comms
#   okta-jarvis / okta-unicorn / unicorn → Jarvis / Unicorn
TEAM_DISPLAY = {
    # Biz
    "biz-enterprise": "Biz",
    "biz": "Biz",
    "biz-admins": "Biz Admins",
    "biz-payments": "Payments",
    # Odyssey (discovery group is Odyssey)
    "discovery": "Odyssey",
    "okta-discovery": "Odyssey",
    "odyssey": "Odyssey",
    # Comms
    "okta-communication": "Comms",
    # Jarvis
    "okta-jarvis": "Jarvis",
    "jarvis": "Jarvis",
    # Unicorn
    "okta-unicorn": "Unicorn",
    "unicorn": "Unicorn",
    # Other squads
    "team_journey": "Journey",
    "okta-airbook": "Airbook",
    "airbook": "Airbook",
    "okta-payments": "Payments",
    "okta-cross": "Cross",
    "cross": "Cross",
    "okta-groot": "Groot",
    "groot": "Groot",
    "okta-starlite": "Starlite",
    "okta-reservo": "Reservo",
    "okta-chronos": "Chronos",
    "yourporter": "YourPorter",
    "okta-pop": "Pop",
    "okta-nexus": "Nexus",
    "demand-tlv": "Demand TLV",
    "sketch": "Sketch",
    "okta-sketch": "Sketch",
    "okta-developers": "Developers",
    "adventure": "Adventure",
    "oxp": "OXP",
    "okta-the-club": "The Club",
    "sasquatch": "Sasquatch",
    "okta-payday": "Payday",
    "property-team": "Property",
    "blitz": "Blitz",
    "okta-gx": "GX",
    "okta-cortex": "Cortex",
    "maximmovchan1": "maximmovchan1",
    "arambse": "arambse",
    "sergey-guesty": "sergey-guesty",
}

INDIVIDUAL_HANDLES = {"arambse", "sergey-guesty", "maximmovchan1"}


def parse_owners(raw: str) -> list[str]:
    if not raw:
        return []
    owners: list[str] = []
    seen: set[str] = set()
    for part in re.split(r"[\n,;]+", raw):
        part = part.strip()
        if not part:
            continue
        handle = part.lstrip("@")
        slug = handle.split("/", 1)[1] if handle.startswith("guestyorg/") else handle
        if slug not in seen:
            seen.add(slug)
            owners.append(slug)
    return owners


def display_name(slug: str) -> str:
    if slug in TEAM_DISPLAY:
        return TEAM_DISPLAY[slug]
    name = slug
    for prefix in ("okta-", "team_", "biz-"):
        if name.startswith(prefix):
            name = name[len(prefix) :]
            break
    return name.replace("-", " ").replace("_", " ").title()


def owner_handle(slug: str) -> str:
    if slug in INDIVIDUAL_HANDLES:
        return f"@{slug}"
    return f"@guestyorg/{slug}"


def sheet_tab_name(name: str) -> str:
    return re.sub(r"[\\/?*\[\]:]", "-", name)[:31] or "Unknown"


def main() -> None:
    with SRC.open(newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))

    enriched = []
    for r in rows:
        owners = parse_owners(r.get("Owners") or "")
        team_names: list[str] = []
        seen_names: set[str] = set()
        for o in owners:
            n = display_name(o)
            if n not in seen_names:
                seen_names.add(n)
                team_names.append(n)
        enriched.append(
            {
                "Team Name": ", ".join(team_names) if team_names else "(unassigned)",
                "Repo": r.get("Repo") or "",
                "Owners": ", ".join(owner_handle(o) for o in owners),
                "Team Keys": team_names or ["(unassigned)"],
                "Owner Slugs": owners,
                "Rule Type": r.get("Rule Type") or "",
                "Severity": r.get("Severity") or "",
                "Message": r.get("Message") or "",
                "Matched Files": r.get("Matched Files") or "",
            }
        )

    by_team: dict[str, list] = defaultdict(list)
    team_github: dict[str, set[str]] = defaultdict(set)
    for row in enriched:
        for name in row["Team Keys"]:
            by_team[name].append(row)
        for slug in row["Owner Slugs"]:
            team_github[display_name(slug)].add(owner_handle(slug))

    summary_rows = []
    for name, team_rows in by_team.items():
        repos = {r["Repo"] for r in team_rows}
        sev = Counter(r["Severity"] for r in team_rows)
        summary_rows.append(
            {
                "Team Name": name,
                "GitHub Owner(s)": ", ".join(sorted(team_github.get(name, []))),
                "Repos": len(repos),
                "Violations": len(team_rows),
                "Errors": sev.get("Error", 0),
                "Warnings": sev.get("Warning", 0),
            }
        )
    summary_rows.sort(key=lambda x: (-x["Violations"], x["Team Name"].lower()))

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    error_fill = PatternFill("solid", fgColor="FCE4D6")
    warn_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")

    def write_table(ws, headers, data_rows, widths, severity_col=None, start_row=1):
        for col, h in enumerate(headers, 1):
            cell = ws.cell(start_row, col, h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for item in data_rows:
            values = [item.get(h, "") for h in headers]
            ws.append(values)
            ridx = ws.max_row
            for c in range(1, len(headers) + 1):
                cell = ws.cell(ridx, c)
                cell.alignment = wrap
                cell.border = thin
            if severity_col is not None:
                sev = values[severity_col]
                fill = (
                    error_fill
                    if sev == "Error"
                    else (warn_fill if sev == "Warning" else None)
                )
                if fill:
                    for c in range(1, len(headers) + 1):
                        ws.cell(ridx, c).fill = fill
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.auto_filter.ref = (
            f"A{start_row}:{get_column_letter(len(headers))}{ws.max_row}"
        )
        ws.freeze_panes = f"A{start_row + 1}"
        ws.row_dimensions[start_row].height = 22

    # Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = (
        "Frontend Governance — Data Rule Violations by Team (2026-07-26)"
    )
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws.merge_cells("A1:F1")
    ws["A2"] = (
        "Owners (column B) mapped to Guesty team names. "
        "Example: biz-enterprise→Biz, discovery→Odyssey, okta-communication→Comms, "
        "okta-jarvis→Jarvis, okta-unicorn/unicorn→Unicorn. "
        "Multi-owner repos appear on each owning team tab."
    )
    ws["A2"].font = Font(italic=True, size=10, color="666666")
    ws.merge_cells("A2:F2")

    headers = [
        "Team Name",
        "GitHub Owner(s)",
        "Repos",
        "Violations",
        "Errors",
        "Warnings",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(4, col, h)
        cell.font = header_font
        cell.fill = header_fill
    for i, item in enumerate(summary_rows):
        r = 5 + i
        for c, h in enumerate(headers, 1):
            cell = ws.cell(r, c, item[h])
            cell.border = thin
            cell.alignment = wrap
    ws.auto_filter.ref = f"A4:F{4 + len(summary_rows)}"
    ws.freeze_panes = "A5"
    for i, w in enumerate([14, 55, 10, 12, 10, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    total_row = 5 + len(summary_rows)
    ws.cell(total_row, 1, "TOTAL unique violation rows").font = Font(bold=True)
    ws.cell(total_row, 3, len({r["Repo"] for r in enriched})).font = Font(bold=True)
    ws.cell(total_row, 4, len(enriched)).font = Font(bold=True)
    ws.cell(
        total_row, 5, sum(1 for r in enriched if r["Severity"] == "Error")
    ).font = Font(bold=True)
    ws.cell(
        total_row, 6, sum(1 for r in enriched if r["Severity"] == "Warning")
    ).font = Font(bold=True)
    ws.cell(
        total_row + 1,
        1,
        "Note: Sum of per-team Violations can exceed TOTAL because "
        "multi-owner repos are counted on each team tab.",
    ).font = Font(italic=True, size=9, color="666666")
    ws.merge_cells(
        start_row=total_row + 1, start_column=1, end_row=total_row + 1, end_column=6
    )

    ws.cell(
        total_row + 3, 1, "Get a shareable Google Sheet link (30 seconds)"
    ).font = Font(bold=True, size=12, color="1F4E79")
    for i, s in enumerate(
        [
            "1. Download Frontend_Governance_Violations_By_Team.xlsx",
            "2. Go to drive.google.com → New → File upload → select this file",
            "3. Right-click → Open with → Google Sheets (keeps all team tabs)",
            "4. Share → Guesty (guesty.com) or Anyone with the link → Copy link",
            "5. Send teams the link; each squad uses their named tab",
        ]
    ):
        ws.cell(total_row + 4 + i, 1, s)

    # Owner → Team mapping reference tab
    ws_map = wb.create_sheet("Owner→Team Map", 1)
    map_headers = ["GitHub Owner", "Team Name"]
    map_rows = [
        {"GitHub Owner": owner_handle(slug), "Team Name": display_name(slug)}
        for slug in sorted(TEAM_DISPLAY.keys(), key=lambda s: (display_name(s), s))
    ]
    # also include any unmapped slugs found in data
    all_slugs = sorted({s for r in enriched for s in r["Owner Slugs"]})
    for slug in all_slugs:
        if slug not in TEAM_DISPLAY:
            map_rows.append(
                {"GitHub Owner": owner_handle(slug), "Team Name": display_name(slug)}
            )
    write_table(ws_map, map_headers, map_rows, [36, 16])

    all_headers = [
        "Team Name",
        "Repo",
        "Owners",
        "Rule Type",
        "Severity",
        "Message",
        "Matched Files",
    ]
    ws_all = wb.create_sheet("All Violations")
    write_table(
        ws_all,
        all_headers,
        [{k: r[k] for k in all_headers} for r in enriched],
        [28, 28, 50, 14, 10, 55, 40],
        severity_col=4,
    )

    used = {"Summary", "Owner→Team Map", "All Violations"}
    for name, _ in sorted(
        by_team.items(), key=lambda kv: (-len(kv[1]), kv[0].lower())
    ):
        title = sheet_tab_name(name)
        assert title not in used, title
        used.add(title)
        ws_t = wb.create_sheet(title)
        team_rows = [
            {
                "Team Name": name,
                "Repo": r["Repo"],
                "Owners": r["Owners"],
                "Rule Type": r["Rule Type"],
                "Severity": r["Severity"],
                "Message": r["Message"],
                "Matched Files": r["Matched Files"],
            }
            for r in by_team[name]
        ]
        write_table(
            ws_t,
            all_headers,
            team_rows,
            [14, 28, 50, 14, 10, 55, 40],
            severity_col=4,
        )

    # Example multi-owner sanity check (user's example)
    example_owners = [
        "biz-enterprise",
        "discovery",
        "okta-communication",
        "okta-jarvis",
        "okta-unicorn",
        "unicorn",
    ]
    example_teams = [display_name(s) for s in example_owners]
    assert example_teams == [
        "Biz",
        "Odyssey",
        "Comms",
        "Jarvis",
        "Unicorn",
        "Unicorn",
    ], example_teams

    outs = [
        Path("/workspace/exports/Frontend_Governance_Violations_By_Team.xlsx"),
        Path("/opt/cursor/artifacts/Frontend_Governance_Violations_By_Team.xlsx"),
        Path("/workspace/frontend-governance/Frontend_Governance_Violations_By_Team.xlsx"),
    ]
    for p in outs:
        p.parent.mkdir(parents=True, exist_ok=True)
        wb.save(p)
        print("saved", p, p.stat().st_size)

    index = Path("/workspace/frontend-governance/team_tabs_index.csv")
    with index.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=headers)
        w.writeheader()
        w.writerows(summary_rows)

    mapping_csv = Path("/workspace/frontend-governance/owner_to_team_map.csv")
    with mapping_csv.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["GitHub Owner", "Team Name"])
        w.writeheader()
        w.writerows(map_rows)

    readme = Path("/workspace/frontend-governance/README.md")
    readme.write_text(
        f"""# Frontend Governance — Data Rule Violations by Team

Source: `Frontend Governance (2026-07-26) - Data Rule Violations` CSV

## Team name mapping (Owners → Team)

| GitHub Owner | Team Name |
|---|---|
| `@guestyorg/biz-enterprise` | **Biz** |
| `@guestyorg/discovery` | **Odyssey** |
| `@guestyorg/okta-communication` | **Comms** |
| `@guestyorg/okta-jarvis` | **Jarvis** |
| `@guestyorg/okta-unicorn` | **Unicorn** |
| `@guestyorg/unicorn` | **Unicorn** |

Full map: `owner_to_team_map.csv` and the **Owner→Team Map** sheet tab.

## File
- `Frontend_Governance_Violations_By_Team.xlsx`

## Structure
- **Summary** — team names, GitHub owners, counts
- **Owner→Team Map** — full owner → team reference
- **All Violations** — full dataset with Team Name
- **One tab per team** ({len(by_team)} teams)

## Create a shareable Google Sheet link
1. Upload the `.xlsx` to [Google Drive](https://drive.google.com)
2. Open with **Google Sheets**
3. **Share** → copy link for teams

## Team tabs ({len(summary_rows)})
"""
        + "\n".join(
            f"- **{r['Team Name']}** — {r['Violations']} violations, {r['Repos']} repos "
            f"(`{r['GitHub Owner(s)']}`)"
            for r in summary_rows
        )
        + "\n",
        encoding="utf-8",
    )

    print("teams", len(summary_rows))
    print("tabs", len(wb.sheetnames))
    print(wb.sheetnames)
    print("example mapping OK:", example_teams)


if __name__ == "__main__":
    main()
