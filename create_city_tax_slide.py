#!/usr/bin/env python3
"""Create a professional, Google Slides-compatible City Tax status deck."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN


NAVY = RGBColor(0x0F, 0x2C, 0x59)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF7, 0xF8, 0xFA)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xE2, 0xE8, 0xF0)
ACCENT_BLUE = RGBColor(0x25, 0x63, 0xEB)
ACCENT_TEAL = RGBColor(0x0D, 0x94, 0x88)
ACCENT_AMBER = RGBColor(0xD9, 0x77, 0x06)
ACCENT_RED = RGBColor(0xDC, 0x26, 0x26)
LIGHT_BLUE = RGBColor(0xEF, 0xF6, 0xFF)
LIGHT_TEAL = RGBColor(0xF0, 0xFD, 0xFA)
LIGHT_AMBER = RGBColor(0xFF, 0xFB, 0xEB)
LIGHT_RED = RGBColor(0xFE, 0xF2, 0xF2)
GREEN = RGBColor(0x05, 0x96, 0x69)
LIGHT_GREEN = RGBColor(0xEC, 0xFD, 0xF5)
GUESTY = RGBColor(0x1A, 0x73, 0xE8)  # approximate brand blue for footer mark


def set_run(run, size=12, bold=False, color=SLATE, font_name="Arial"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name


def add_textbox(slide, left, top, width, height, text, size=12, bold=False,
                color=SLATE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_rounded_rect(slide, left, top, width, height, fill, line=None, line_w=Pt(1)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = line_w
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def add_rect(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    return shape


def add_pill(slide, left, top, width, height, fill, text, text_color=WHITE, size=10, border=None):
    shape = add_rounded_rect(slide, left, top, width, height, fill, line=border, line_w=Pt(1))
    try:
        shape.adjustments[0] = 0.5
    except Exception:
        pass
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=True, color=text_color)
    return shape


def add_bullet_block(slide, left, top, width, height, items, body_size=12):
    """items: list of (title, body, emphasis) where emphasis in {False, True, 'alert'}"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True

    first = True
    for title, body, emphasis in items:
        if not first:
            sp = tf.add_paragraph()
            sp.text = ""
            sp.space_before = Pt(6)
            sp.space_after = Pt(0)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(2)
        p.space_after = Pt(2)

        bullet = p.add_run()
        bullet.text = "●  "
        set_run(bullet, size=body_size - 1, bold=False, color=ACCENT_BLUE)

        t_run = p.add_run()
        t_run.text = title
        set_run(t_run, size=body_size, bold=True, color=NAVY)

        if body:
            b = tf.add_paragraph()
            b.alignment = PP_ALIGN.LEFT
            b.space_before = Pt(1)
            b.space_after = Pt(4)
            b_run = b.add_run()
            b_run.text = "    " + body
            body_color = ACCENT_RED if emphasis == "alert" else (MUTED if emphasis else SLATE)
            body_bold = True if emphasis == "alert" else False
            set_run(b_run, size=body_size - 1, bold=body_bold, color=body_color)
    return box


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ── Slide 1: City Tax Challenges & Next Steps ─────────────────────────
    slide = prs.slides.add_slide(blank)
    add_rect(slide, Inches(0), Inches(0), prs.slide_width, prs.slide_height, BG)
    add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(0.08), NAVY)

    add_textbox(
        slide, Inches(0.55), Inches(0.28), Inches(8.5), Inches(0.45),
        "City Tax — Challenges & Next Steps",
        size=26, bold=True, color=NAVY
    )
    add_pill(
        slide, Inches(11.15), Inches(0.35), Inches(1.55), Inches(0.34),
        ACCENT_AMBER, "IN PROGRESS", WHITE, size=10
    )
    add_textbox(
        slide, Inches(0.55), Inches(0.72), Inches(9), Inches(0.28),
        "Status update  ·  EU regulatory readiness & product priorities",
        size=12, bold=False, color=MUTED
    )

    # Country priority banner
    add_rounded_rect(
        slide, Inches(0.45), Inches(1.12), Inches(12.4), Inches(0.95),
        LIGHT_BLUE, line=RGBColor(0xBF, 0xDB, 0xFE), line_w=Pt(1.25)
    )
    add_rect(slide, Inches(0.45), Inches(1.12), Inches(0.12), Inches(0.95), ACCENT_BLUE)

    add_textbox(
        slide, Inches(0.75), Inches(1.18), Inches(4), Inches(0.28),
        "COUNTRY PRIORITY",
        size=10, bold=True, color=ACCENT_BLUE
    )

    # Priority chips
    chips = [
        (Inches(0.75), ACCENT_RED, LIGHT_RED, "Italy · Critical"),
        (Inches(3.0), ACCENT_AMBER, LIGHT_AMBER, "Spain · High"),
        (Inches(5.15), ACCENT_AMBER, LIGHT_AMBER, "France · High"),
    ]
    for x, edge, fill, label in chips:
        add_pill(slide, x, Inches(1.52), Inches(2.05), Inches(0.30), fill, label, edge, size=10, border=edge)

    add_textbox(
        slide, Inches(7.5), Inches(1.48), Inches(5.0), Inches(0.40),
        "Focus markets for dynamic PMC city-tax setup",
        size=11, bold=False, color=SLATE
    )

    # Three content cards
    card_top = Inches(2.3)
    card_h = Inches(4.55)
    card_w = Inches(3.95)
    gap = Inches(0.2)
    starts = [Inches(0.45), Inches(0.45) + card_w + gap, Inches(0.45) + 2 * (card_w + gap)]

    cards = [
        {
            "accent": ACCENT_RED,
            "header_bg": LIGHT_RED,
            "eyebrow": "01  CURRENT STATE",
            "title": "Gaps & Challenges",
            "items": [
                (
                    "Fragmented EU Rules",
                    "Each EU country has its own city tax requirements — no single shared model.",
                    False,
                ),
                (
                    "Product Capability Gap",
                    "Product must encode tax regulation per priority country and allow dynamic setup for PMCs.",
                    False,
                ),
            ],
        },
        {
            "accent": ACCENT_TEAL,
            "header_bg": LIGHT_TEAL,
            "eyebrow": "02  ACTIONS",
            "title": "Next Steps",
            "items": [
                (
                    "EU Tax Advisor Partnership",
                    "Partnership is closing with an EU tax advisor to map tax requirements across priority countries.",
                    False,
                ),
                (
                    "Client Discovery",
                    "Product will interview EU clients to learn how they currently handle city tax setup.",
                    False,
                ),
            ],
        },
        {
            "accent": ACCENT_AMBER,
            "header_bg": LIGHT_AMBER,
            "eyebrow": "03  DECISIONS",
            "title": "Going Forward",
            "items": [
                (
                    "Roadmap Priority — Q4",
                    "Initiative is currently a Q4 candidate.",
                    False,
                ),
                (
                    "Unblock Priority Markets",
                    "Target: unlock Italy, Spain & France readiness in Q1 2027.",
                    False,
                ),
            ],
        },
    ]

    for left, card in zip(starts, cards):
        add_rounded_rect(slide, left, card_top, card_w, card_h, CARD, line=BORDER, line_w=Pt(1.25))
        add_rect(slide, left, card_top, card_w, Inches(0.12), card["accent"])
        add_rect(slide, left, card_top + Inches(0.12), card_w, Inches(0.72), card["header_bg"])

        add_textbox(
            slide, left + Inches(0.2), card_top + Inches(0.18), card_w - Inches(0.35), Inches(0.22),
            card["eyebrow"], size=9, bold=True, color=card["accent"]
        )
        add_textbox(
            slide, left + Inches(0.2), card_top + Inches(0.4), card_w - Inches(0.35), Inches(0.35),
            card["title"], size=16, bold=True, color=NAVY
        )
        add_bullet_block(
            slide,
            left + Inches(0.18),
            card_top + Inches(1.0),
            card_w - Inches(0.35),
            card_h - Inches(1.2),
            card["items"],
            body_size=12,
        )

    # Footer
    add_rect(slide, Inches(0), Inches(7.15), prs.slide_width, Inches(0.35), NAVY)
    add_textbox(
        slide, Inches(0.55), Inches(7.18), Inches(8), Inches(0.28),
        "Guesty  ·  Confidential  ·  Product & Compliance  ·  Editable Google Slides / PowerPoint",
        size=10, bold=False, color=RGBColor(0xCB, 0xD5, 0xE1)
    )
    add_textbox(
        slide, Inches(10.5), Inches(7.18), Inches(2.3), Inches(0.28),
        "Slide 1 of 2",
        size=10, bold=False, color=RGBColor(0xCB, 0xD5, 0xE1), align=PP_ALIGN.RIGHT
    )

    # ── Slide 2: Research & Decision Framework ────────────────────────────
    slide2 = prs.slides.add_slide(blank)
    add_rect(slide2, Inches(0), Inches(0), prs.slide_width, prs.slide_height, BG)
    add_rect(slide2, Inches(0), Inches(0), prs.slide_width, Inches(0.08), NAVY)

    add_textbox(
        slide2, Inches(0.55), Inches(0.28), Inches(10), Inches(0.45),
        "City Tax — Research & Decision Framework",
        size=26, bold=True, color=NAVY
    )
    add_textbox(
        slide2, Inches(0.55), Inches(0.72), Inches(10), Inches(0.28),
        "Working session slide — capture owners, findings, and Q4 commitment",
        size=12, bold=False, color=MUTED
    )

    # Left panel: Research workstreams
    add_rounded_rect(
        slide2, Inches(0.45), Inches(1.2), Inches(6.05), Inches(5.55),
        CARD, line=BORDER, line_w=Pt(1.25)
    )
    add_rect(slide2, Inches(0.45), Inches(1.2), Inches(6.05), Inches(0.12), ACCENT_TEAL)
    add_textbox(
        slide2, Inches(0.7), Inches(1.5), Inches(5.5), Inches(0.35),
        "Research Workstreams", size=18, bold=True, color=NAVY
    )

    research_items = [
        (
            "A. Tax Advisor Mapping",
            "Close partnership with an EU tax advisor to map city-tax requirements for Italy, Spain, and France.",
            False,
        ),
        (
            "B. Client Interviews",
            "Product interviews with EU clients on how they currently configure and operate city tax.",
            False,
        ),
        (
            "C. Product Implications",
            "Translate findings into regulation-per-country config + dynamic PMC setup requirements.",
            False,
        ),
    ]
    add_bullet_block(slide2, Inches(0.7), Inches(2.05), Inches(5.5), Inches(4.4), research_items, body_size=13)

    # Right panel: Timeline & decisions
    add_rounded_rect(
        slide2, Inches(6.75), Inches(1.2), Inches(6.1), Inches(5.55),
        CARD, line=BORDER, line_w=Pt(1.25)
    )
    add_rect(slide2, Inches(6.75), Inches(1.2), Inches(6.1), Inches(0.12), ACCENT_AMBER)
    add_textbox(
        slide2, Inches(7.0), Inches(1.5), Inches(5.6), Inches(0.35),
        "Timeline & Open Decisions", size=18, bold=True, color=NAVY
    )

    decision_items = [
        (
            "Q4 Candidate Initiative",
            "Currently proposed as a Q4 roadmap priority.",
            False,
        ),
        (
            "Q1 2027 Market Unblock",
            "Goal: enable priority countries (Italy critical; Spain & France high) in Q1 2027.",
            False,
        ),
        (
            "Open Decision",
            "Confirm Q4 commitment, owners, and success criteria before locking roadmap.",
            "alert",
        ),
        (
            "Dependency Note",
            "May interact with broader tax / e-invoicing workstreams — align sequencing.",
            False,
        ),
    ]
    add_bullet_block(slide2, Inches(7.0), Inches(2.05), Inches(5.55), Inches(4.4), decision_items, body_size=12)

    add_rect(slide2, Inches(0), Inches(7.15), prs.slide_width, Inches(0.35), NAVY)
    add_textbox(
        slide2, Inches(0.55), Inches(7.18), Inches(8), Inches(0.28),
        "Guesty  ·  Confidential  ·  Product & Compliance  ·  Editable Google Slides / PowerPoint",
        size=10, bold=False, color=RGBColor(0xCB, 0xD5, 0xE1)
    )
    add_textbox(
        slide2, Inches(10.5), Inches(7.18), Inches(2.3), Inches(0.28),
        "Slide 2 of 2",
        size=10, bold=False, color=RGBColor(0xCB, 0xD5, 0xE1), align=PP_ALIGN.RIGHT
    )

    out_art = "/opt/cursor/artifacts/City_Tax_Challenges_Next_Steps.pptx"
    out_ws = "/workspace/City_Tax_Challenges_Next_Steps.pptx"
    prs.save(out_art)
    prs.save(out_ws)
    print(out_art)
    return out_art


if __name__ == "__main__":
    build()
