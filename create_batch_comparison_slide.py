#!/usr/bin/env python3
"""Google Slides–importable PowerPoint for Accounts Comparison Per Batch chart."""

import os
import subprocess
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from create_batch_comparison_chart import (
    DOMAINS,
    MAY_BATCH,
    JUNE_BATCH,
    OUTPUT_PNG,
    create_chart,
)

COLOR_BG = RGBColor(245, 240, 232)
COLOR_TEXT = RGBColor(30, 41, 59)
COLOR_MAY = RGBColor(45, 75, 30)
COLOR_JUNE = RGBColor(209, 226, 245)
COLOR_JULY = RGBColor(180, 83, 9)
COLOR_HEADER = RGBColor(45, 75, 30)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_MID = RGBColor(100, 116, 139)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "Arial"
OUTPUT_PPTX = "/workspace/accounts_comparison_per_batch.pptx"


def set_slide_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG


def add_textbox(slide, left, top, width, height, text, size=12, bold=False,
                color=COLOR_TEXT, align=PP_ALIGN.LEFT, italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.italic = italic
    p.font.color.rgb = color
    p.alignment = align
    return box


def build_chart_slide(slide):
    set_slide_bg(slide)
    if os.path.exists(OUTPUT_PNG):
        slide.shapes.add_picture(OUTPUT_PNG, Inches(0.55), Inches(0.35), width=Inches(12.2))
    else:
        add_textbox(slide, Inches(2), Inches(3), Inches(9), Inches(1),
                    "Chart image missing. Run: python3 create_batch_comparison_chart.py",
                    size=14, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(0.6), Inches(6.95), Inches(12), Inches(0.35),
                "Slide 2 contains editable data · Upload to Google Drive → Open with Google Slides",
                size=9, color=COLOR_MID, italic=True, align=PP_ALIGN.CENTER)


def build_data_slide(slide):
    set_slide_bg(slide)
    add_textbox(slide, Inches(0.6), Inches(0.4), Inches(12), Inches(0.5),
                "Accounts Comparison Per Batch — Data (Editable)", size=24, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(0.6), Inches(0.95), Inches(12), Inches(0.3),
                "Edit values below, then regenerate the chart if needed.",
                size=11, color=COLOR_MID, italic=True, align=PP_ALIGN.CENTER)

    headers = ["Domain", "May Batch Count", "June Batch Count", "July", "May+Jun Total"]
    rows = []
    for domain, may, june in zip(DOMAINS, MAY_BATCH, JUNE_BATCH):
        rows.append([domain, str(may), str(june), "GA (all)", str(may + june)])

    table_shape = slide.shapes.add_table(
        len(rows) + 1, len(headers), Inches(0.7), Inches(1.5), Inches(11.9), Inches(4.6))
    table = table_shape.table
    col_widths = [Inches(2.2), Inches(2.4), Inches(2.4), Inches(2.2), Inches(2.2)]
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    header_fmt = lambda cell, text: _style_cell(cell, text, header=True)
    for ci, h in enumerate(headers):
        header_fmt(table.cell(0, ci), h)

    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            bold = ci == 0
            _style_cell(table.cell(ri, ci), val, bold=bold)

    # Legend note
    add_textbox(slide, Inches(0.7), Inches(6.35), Inches(12), Inches(0.45),
                "■ May Batch Count (dark green)    ■ June Batch Count (light blue)    "
                "■ July — GA for all accounts (amber; no per-domain count)",
                size=10, color=COLOR_TEXT)


def _style_cell(cell, text, header=False, bold=False):
    cell.text = text
    p = cell.text_frame.paragraphs[0]
    p.font.name = FONT
    p.font.size = Pt(11 if header else 10)
    p.font.bold = header or bold
    if header:
        p.font.color.rgb = COLOR_WHITE
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_HEADER
    else:
        p.font.color.rgb = COLOR_TEXT


def build_deck(output_path=OUTPUT_PPTX):
    if not os.path.exists(OUTPUT_PNG):
        create_chart()

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    s_chart = prs.slides.add_slide(blank)
    s_data = prs.slides.add_slide(blank)
    build_chart_slide(s_chart)
    build_data_slide(s_data)

    prs.save(output_path)
    print(f"Saved: {output_path}")
    print("Slides:")
    print("  1. Accounts Comparison Per Batch chart")
    print("  2. Editable data table")
    print("\nOpen in Google Slides:")
    print("  1. Upload .pptx to Google Drive")
    print("  2. Right-click → Open with → Google Slides")


if __name__ == "__main__":
    build_deck()
