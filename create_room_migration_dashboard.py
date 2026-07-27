#!/usr/bin/env python3
"""
Build Room Migration Progress chart from accounts + Guesty Segment column.

Expects exports/accounts_created_with_segment.csv produced by:
  python3 enrich_segments_from_guesty.py

That file must contain: Account ID, Created At, Segment
"""

from __future__ import annotations

import csv
import json
import re
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
CONFIG_PATH = ROOT / "room_migration_config.json"

BG = "#F5F1E9"
NAVY = "#1B3A5F"
CARD = "#D9EAF7"
CHART_BG = "#FFFFFF"
CHART_BORDER = "#D0D5DD"
MUTED = "#6B7C93"
GRID = "#E6EAF0"

SEGMENT_ORDER = ("SMB", "Mid-Market", "SME", "Enterprise", "Unknown")
SEGMENT_COLORS = {
    "SMB": "#4A8FD4",
    "Mid-Market": "#E35D5B",
    "SME": "#F0C04A",
    "Enterprise": "#7C5CBF",
    "Unknown": "#9AA5B5",
}
SEGMENT_ALIASES = {
    "smb": "SMB",
    "mid-market": "Mid-Market",
    "mid market": "Mid-Market",
    "midmarket": "Mid-Market",
    "sme": "SME",
    "enterprise": "Enterprise",
}


def load_config() -> dict:
    with open(CONFIG_PATH, encoding="utf-8") as f:
        return json.load(f)


def parse_date(value: str) -> datetime | None:
    if not value or value.strip().lower() in {"null", "none", "nan", ""}:
        return None
    value = value.strip()
    if "T" in value or value.endswith("Z"):
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).replace(tzinfo=None)
        except ValueError:
            pass
    for fmt in ("%b %d, %Y", "%B %d, %Y", "%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    m = re.match(r"^(\d{4}-\d{2}-\d{2})", value)
    if m:
        return datetime.strptime(m.group(1), "%Y-%m-%d")
    return None


def monday_of(dt: datetime) -> datetime:
    return (dt - timedelta(days=dt.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)


def normalize_segment(raw: str | None) -> str:
    if raw is None:
        return "Unknown"
    value = str(raw).strip()
    if not value or value.lower() in {"null", "none", "nan"}:
        return "Unknown"
    return SEGMENT_ALIASES.get(value.lower(), value)


def load_enriched_accounts(path: Path) -> list[dict]:
    if not path.exists():
        raise SystemExit(
            f"Missing {path}.\n"
            "First enrich Segments from Guesty:\n"
            "  python3 enrich_segments_from_guesty.py\n"
            "See ROOM_MIGRATION_README.md for Guesty HQ auth / export instructions."
        )
    rows = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            created = parse_date(row.get("Created At") or row.get("createdAt") or "")
            if not created:
                continue
            rows.append(
                {
                    "Account ID": (row.get("Account ID") or row.get("accountId") or "").strip(),
                    "Created At": created,
                    "Segment": normalize_segment(row.get("Segment")),
                    "Segment Source": row.get("Segment Source", ""),
                }
            )
    rows.sort(key=lambda r: r["Created At"])
    return rows


def aggregate(rows: list[dict], chart_start: datetime) -> dict:
    totals = {seg: 0 for seg in SEGMENT_ORDER}
    weekly_new = defaultdict(lambda: {seg: 0 for seg in SEGMENT_ORDER})

    for row in rows:
        seg = row["Segment"] if row["Segment"] in totals else "Unknown"
        if seg not in totals:
            totals[seg] = 0
        totals[seg] += 1
        week = monday_of(row["Created At"])
        if seg not in weekly_new[week]:
            weekly_new[week][seg] = 0
        weekly_new[week][seg] += 1

    extras = [s for s in totals if s not in SEGMENT_ORDER]
    segment_order = list(SEGMENT_ORDER) + sorted(extras, key=lambda s: (-totals[s], s))

    if rows:
        end = monday_of(max(r["Created At"] for r in rows))
        latest = max(r["Created At"] for r in rows)
        earliest = min(r["Created At"] for r in rows)
    else:
        end = monday_of(datetime.today())
        latest = earliest = datetime.today()

    start = monday_of(chart_start)
    if start > end:
        start = end

    weeks: list[datetime] = []
    cursor = start
    while cursor <= end:
        weeks.append(cursor)
        cursor += timedelta(days=7)

    cumulative = {seg: [] for seg in segment_order}
    running = {seg: 0 for seg in segment_order}
    for week in weeks:
        for seg in segment_order:
            running[seg] += weekly_new[week].get(seg, 0)
            cumulative[seg].append(running[seg])

    known = sum(v for k, v in totals.items() if k != "Unknown")
    return {
        "totals": totals,
        "total": len(rows),
        "known_segment": known,
        "unknown_segment": totals.get("Unknown", 0),
        "segment_order": segment_order,
        "weeks": weeks,
        "cumulative": cumulative,
        "auto_week_label": f"{earliest.strftime('%B %-d')} - {latest.strftime('%B %-d')}",
    }


def fmt_int(n: int) -> str:
    return f"{n:,}"


def write_segment_summary(agg: dict, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Segment", "Accounts"])
        for seg in agg["segment_order"]:
            writer.writerow([seg, agg["totals"].get(seg, 0)])
        writer.writerow(["TOTAL", agg["total"]])


def segment_color(seg: str) -> str:
    return SEGMENT_COLORS.get(seg, "#7C8A9A")


def draw_dashboard_image(agg: dict, week_label: str, title: str, total_label: str, out_path: Path) -> None:
    fig = plt.figure(figsize=(13.333, 7.5), dpi=150, facecolor=BG)
    fig.subplots_adjust(left=0.05, right=0.97, top=0.92, bottom=0.08)
    fig.text(0.05, 0.92, title, fontsize=28, fontweight="bold", color=NAVY, ha="left", va="top")

    ax_date = fig.add_axes([0.05, 0.82, 0.26, 0.055])
    ax_date.set_xlim(0, 1)
    ax_date.set_ylim(0, 1)
    ax_date.axis("off")
    ax_date.add_patch(
        FancyBboxPatch((0.0, 0.05), 0.98, 0.85, boxstyle="round,pad=0.02,rounding_size=0.2", linewidth=0, facecolor="#E4E7EC")
    )
    ax_date.text(0.5, 0.48, week_label, ha="center", va="center", fontsize=11, color=NAVY)

    def metric_card(bounds, label, value, sub=None, value_size=26):
        ax = fig.add_axes(bounds)
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 1)
        ax.axis("off")
        ax.add_patch(
            FancyBboxPatch((0.02, 0.08), 0.96, 0.84, boxstyle="round,pad=0.02,rounding_size=0.15", linewidth=0, facecolor=CARD)
        )
        ax.text(0.5, 0.72, label, ha="center", va="center", fontsize=11, color=NAVY)
        ax.text(0.5, 0.38 if sub else 0.35, value, ha="center", va="center", fontsize=value_size, fontweight="bold", color=NAVY)
        if sub:
            ax.text(0.5, 0.14, sub, ha="center", va="center", fontsize=8, color=MUTED)

    metric_card([0.35, 0.78, 0.30, 0.12], total_label, fmt_int(agg["total"]), value_size=28)

    card_segs = [s for s in agg["segment_order"] if agg["totals"].get(s, 0) > 0 or s in ("SMB", "Mid-Market", "SME", "Unknown")]
    n = max(len(card_segs), 1)
    width = min(0.22, 0.88 / n)
    gap = (0.90 - n * width) / (n + 1)
    for i, seg in enumerate(card_segs):
        left = 0.05 + gap + i * (width + gap)
        sub = f"{fmt_int(agg['known_segment'])} known from Guesty" if seg == "Unknown" else None
        metric_card([left, 0.62, width, 0.12], seg, fmt_int(agg["totals"].get(seg, 0)), sub=sub, value_size=24)

    ax = fig.add_axes([0.08, 0.10, 0.86, 0.46])
    for spine in ax.spines.values():
        spine.set_color(CHART_BORDER)
    ax.set_facecolor(CHART_BG)
    ax.set_title("Cumulative Trends by Segment", loc="left", fontsize=12, color=MUTED, pad=10)

    weeks = agg["weeks"]
    x = list(range(len(weeks)))
    width_bar = 0.55 if len(weeks) > 8 else 0.65
    bottom = [0] * len(weeks)
    for seg in agg["segment_order"]:
        vals = agg["cumulative"][seg]
        ax.bar(x, vals, width=width_bar, bottom=bottom, color=segment_color(seg), label=seg, zorder=3)
        bottom = [b + v for b, v in zip(bottom, vals)]

    ax.set_ylabel("Count", color=MUTED, fontsize=10)
    ax.set_xlabel("Week starting", color=MUTED, fontsize=10)
    ax.set_xticks(x)
    ax.set_xticklabels([w.strftime("%-m/%d/%Y") for w in weeks], rotation=45, ha="right", fontsize=9, color=MUTED)
    ax.tick_params(axis="y", colors=MUTED, labelsize=9)
    ax.yaxis.grid(True, color=GRID, linewidth=0.8, zorder=0)
    ax.set_axisbelow(True)
    ax.set_ylim(0, max(10, int(max(agg["total"], 1) * 1.12)))
    ax.legend(loc="upper center", bbox_to_anchor=(0.55, 1.02), ncol=min(5, len(agg["segment_order"])), frameon=False, fontsize=9)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, facecolor=fig.get_facecolor(), bbox_inches="tight", pad_inches=0.25)
    plt.close(fig)


def draw_chart_only(agg: dict, out_path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12.2, 4.0), dpi=150, facecolor=CHART_BG)
    ax.set_facecolor(CHART_BG)
    weeks = agg["weeks"]
    x = list(range(len(weeks)))
    width_bar = 0.55 if len(weeks) > 8 else 0.65
    bottom = [0] * len(weeks)
    for seg in agg["segment_order"]:
        vals = agg["cumulative"][seg]
        ax.bar(x, vals, width=width_bar, bottom=bottom, color=segment_color(seg), label=seg, zorder=3)
        bottom = [b + v for b, v in zip(bottom, vals)]
    ax.set_title("Cumulative Trends by Segment", loc="left", fontsize=12, color=MUTED, pad=12)
    ax.set_ylabel("Count", color=MUTED, fontsize=10)
    ax.set_xlabel("Week starting", color=MUTED, fontsize=10)
    ax.set_xticks(x)
    ax.set_xticklabels([w.strftime("%-m/%d/%Y") for w in weeks], rotation=45, ha="right", fontsize=9, color=MUTED)
    ax.tick_params(axis="y", colors=MUTED, labelsize=9)
    ax.yaxis.grid(True, color=GRID, linewidth=0.8, zorder=0)
    ax.set_axisbelow(True)
    for spine in ax.spines.values():
        spine.set_color(CHART_BORDER)
    ax.set_ylim(0, max(10, int(max(agg["total"], 1) * 1.12)))
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, 1.02), ncol=min(5, len(agg["segment_order"])), frameon=False, fontsize=9)
    fig.tight_layout()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, facecolor=CHART_BG, bbox_inches="tight", pad_inches=0.15)
    plt.close(fig)


def set_run(run, size=12, bold=False, color=RGBColor(0x1B, 0x3A, 0x5F), font_name="Arial"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name


def add_textbox(slide, left, top, width, height, text, size=12, bold=False, color=RGBColor(0x1B, 0x3A, 0x5F), align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_rounded_rect(slide, left, top, width, height, fill: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    try:
        shape.adjustments[0] = 0.12
    except Exception:
        pass
    return shape


def add_metric_card(slide, left, top, width, height, label, value, sub=None, value_size=24):
    card = add_rounded_rect(slide, left, top, width, height, RGBColor(0xD9, 0xEA, 0xF7))
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    set_run(run, size=12)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    run2.text = value
    set_run(run2, size=value_size, bold=True)
    if sub:
        p3 = tf.add_paragraph()
        p3.alignment = PP_ALIGN.CENTER
        run3 = p3.add_run()
        run3.text = sub
        set_run(run3, size=9, color=RGBColor(0x6B, 0x7C, 0x93))
    return card


def build_pptx(agg: dict, week_label: str, title: str, total_label: str, chart_path: Path, out_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0xF5, 0xF1, 0xE9)
    bg.line.fill.background()

    add_textbox(slide, Inches(0.55), Inches(0.28), Inches(7), Inches(0.5), title, size=28, bold=True)
    date_box = add_rounded_rect(slide, Inches(0.55), Inches(0.85), Inches(2.8), Inches(0.38), RGBColor(0xE4, 0xE7, 0xEC))
    tf = date_box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = week_label
    set_run(run, size=12)

    add_metric_card(slide, Inches(4.6), Inches(0.28), Inches(4.0), Inches(1.0), total_label, fmt_int(agg["total"]), value_size=28)

    card_segs = [s for s in agg["segment_order"] if agg["totals"].get(s, 0) > 0 or s in ("SMB", "Mid-Market", "SME", "Unknown")]
    n = max(len(card_segs), 1)
    card_w = min(2.9, (12.2 - 0.2 * (n - 1)) / n)
    for i, seg in enumerate(card_segs):
        left = 0.55 + i * (card_w + 0.2)
        sub = f"{fmt_int(agg['known_segment'])} known from Guesty" if seg == "Unknown" else None
        add_metric_card(slide, Inches(left), Inches(1.5), Inches(card_w), Inches(1.05), seg, fmt_int(agg["totals"].get(seg, 0)), sub=sub)

    if chart_path.exists():
        slide.shapes.add_picture(str(chart_path), Inches(0.55), Inches(2.85), width=Inches(12.2))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)


def build_html(agg: dict, week_label: str, title: str, total_label: str, out_path: Path) -> None:
    weeks = [w.strftime("%-m/%d/%Y") for w in agg["weeks"]]
    datasets = [
        {"label": seg, "data": agg["cumulative"][seg], "backgroundColor": segment_color(seg), "stack": "segments"}
        for seg in agg["segment_order"]
    ]
    cards = "".join(
        f"""<div class="tier-card"><div class="label">{seg}</div><div class="value">{fmt_int(agg['totals'].get(seg, 0))}</div>
        {"<div class='sub'>" + fmt_int(agg['known_segment']) + " known from Guesty</div>" if seg == "Unknown" else ""}</div>"""
        for seg in agg["segment_order"]
        if agg["totals"].get(seg, 0) > 0 or seg in ("SMB", "Mid-Market", "SME", "Unknown")
    )
    summary_rows = "".join(
        f"<tr><td>{seg}</td><td>{fmt_int(agg['totals'].get(seg, 0))}</td></tr>" for seg in agg["segment_order"]
    )
    html = f"""<!DOCTYPE html>
<html lang="en"><head>
<meta charset="UTF-8" /><meta name="viewport" content="width=device-width, initial-scale=1.0" />
<title>{title}</title>
<script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.1/dist/chart.umd.min.js"></script>
<style>
:root {{ --bg:{BG}; --navy:{NAVY}; --card:{CARD}; --muted:{MUTED}; --border:{CHART_BORDER}; }}
* {{ box-sizing:border-box; margin:0; padding:0; }}
body {{ font-family:"Segoe UI","Helvetica Neue",Arial,sans-serif; background:#ddd6c8; color:var(--navy);
  min-height:100vh; padding:1.5rem; display:flex; flex-direction:column; align-items:center; gap:1.25rem; }}
.slide {{ width:min(1280px,100%); aspect-ratio:16/9; background:var(--bg); border-radius:8px;
  box-shadow:0 12px 40px rgba(27,58,95,.12); padding:1.75rem 2rem 1.25rem; display:grid; grid-template-rows:auto auto 1fr; gap:.85rem; }}
.panel {{ width:min(1280px,100%); background:var(--bg); border-radius:8px; box-shadow:0 12px 40px rgba(27,58,95,.12); padding:1.25rem 1.75rem 1.75rem; }}
.header {{ display:grid; grid-template-columns:1fr auto; gap:1rem; }}
h1 {{ font-size:clamp(1.5rem,2.3vw,2rem); font-weight:700; }}
.date-pill {{ display:inline-block; margin-top:.55rem; background:#E4E7EC; border-radius:10px; padding:.35rem .85rem; }}
.total-card,.tier-card {{ background:var(--card); border-radius:14px; text-align:center; padding:.75rem .85rem; box-shadow:0 2px 8px rgba(27,58,95,.08); }}
.tiers {{ display:grid; grid-template-columns:repeat(4,1fr); gap:.75rem; }}
.label {{ font-size:.88rem; }} .value {{ font-size:clamp(1.4rem,2.5vw,2rem); font-weight:700; }}
.sub {{ font-size:.72rem; color:var(--muted); margin-top:.2rem; }}
.chart-panel {{ background:#fff; border:1px solid var(--border); border-radius:10px; padding:.75rem 1rem .4rem; display:flex; flex-direction:column; min-height:0; }}
.chart-title {{ color:var(--muted); font-size:.92rem; }} .chart-wrap {{ flex:1; position:relative; min-height:210px; }}
table {{ width:100%; border-collapse:collapse; margin-top:.75rem; }}
th,td {{ text-align:left; padding:.65rem .75rem; border-bottom:1px solid #e5e7eb; }}
th {{ color:var(--muted); font-size:.78rem; text-transform:uppercase; }}
.note {{ color:var(--muted); font-size:.9rem; margin-top:.4rem; }}
@media (max-width:900px) {{ .slide {{ aspect-ratio:auto; }} .header,.tiers {{ grid-template-columns:1fr; }} }}
</style></head><body>
<div class="slide">
  <div class="header"><div><h1>{title}</h1><div class="date-pill">{week_label}</div></div>
  <div class="total-card"><div class="label">{total_label}</div><div class="value">{fmt_int(agg["total"])}</div></div></div>
  <div class="tiers">{cards}</div>
  <div class="chart-panel"><div class="chart-title">Cumulative Trends by Segment</div><div class="chart-wrap"><canvas id="trendChart"></canvas></div></div>
</div>
<div class="panel">
  <h2>Segment breakdown (from Guesty)</h2>
  <p class="note">Segment is looked up in Guesty per Account ID. Known: {fmt_int(agg["known_segment"])} / {fmt_int(agg["total"])}.</p>
  <table><thead><tr><th>Segment</th><th>Accounts</th></tr></thead><tbody>
  {summary_rows}
  <tr><td><strong>TOTAL</strong></td><td><strong>{fmt_int(agg["total"])}</strong></td></tr>
  </tbody></table>
</div>
<script>
new Chart(document.getElementById("trendChart"), {{
  type:"bar",
  data:{{ labels:{json.dumps(weeks)}, datasets:{json.dumps(datasets)} }},
  options:{{
    responsive:true, maintainAspectRatio:false,
    plugins:{{ legend:{{ position:"top", labels:{{ boxWidth:12, color:"{MUTED}" }} }} }},
    scales:{{
      x:{{ stacked:true, title:{{ display:true, text:"Week starting", color:"{MUTED}" }},
           ticks:{{ maxRotation:45, minRotation:45, color:"{MUTED}" }}, grid:{{ display:false }} }},
      y:{{ stacked:true, beginAtZero:true, suggestedMax:Math.ceil({max(agg["total"], 1)} * 1.12),
           title:{{ display:true, text:"Count", color:"{MUTED}" }},
           ticks:{{ color:"{MUTED}", precision:0 }}, grid:{{ color:"{GRID}" }} }}
    }}
  }}
}});
</script></body></html>"""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(html, encoding="utf-8")


def main() -> None:
    cfg = load_config()
    enriched_path = ROOT / cfg.get("enriched_accounts_csv", "exports/accounts_created_with_segment.csv")
    rows = load_enriched_accounts(enriched_path)
    agg = aggregate(rows, datetime.strptime(cfg["chart_start"], "%Y-%m-%d"))
    week_label = cfg.get("report_week_label") or agg["auto_week_label"]
    title = cfg.get("title", "Room Migration Progress")
    total_label = cfg.get("total_label", "Total Accounts Created")

    segment_csv = ROOT / cfg["output_segment_summary"]
    chart_path = ROOT / cfg["output_chart"]
    html_path = ROOT / cfg["output_html"]
    pptx_path = ROOT / cfg["output_pptx"]
    full_png = ROOT / "exports" / "room-migration-progress.png"

    write_segment_summary(agg, segment_csv)
    draw_chart_only(agg, chart_path)
    draw_dashboard_image(agg, week_label, title, total_label, full_png)
    build_html(agg, week_label, title, total_label, html_path)
    build_pptx(agg, week_label, title, total_label, chart_path, pptx_path)

    print(f"Accounts: {agg['total']}")
    print(f"Known Segment from Guesty: {agg['known_segment']}")
    print("Segment counts:")
    for seg in agg["segment_order"]:
        print(f"  {seg}: {agg['totals'].get(seg, 0)}")
    print(f"Wrote {segment_csv}")
    print(f"Wrote {html_path}")
    print(f"Wrote {pptx_path}")
    print(f"Wrote {full_png}")


if __name__ == "__main__":
    main()
