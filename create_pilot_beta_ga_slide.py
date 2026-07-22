#!/usr/bin/env python3
"""Google Slides–importable PowerPoint for Pilot / Beta / GA accounts chart."""

import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from create_pilot_beta_ga_chart import (
    MONTHS,
    PILOT,
    BETA,
    OUTPUT_PNG,
    create_chart,
)

COLOR_BG = RGBColor(255, 255, 255)
COLOR_TEXT = RGBColor(51, 65, 85)
COLOR_MID = RGBColor(100, 116, 139)
COLOR_HEADER = RGBColor(91, 124, 153)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_GA = RGBColor(194, 65, 12)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "Arial"
OUTPUT_PPTX = "/workspace/pilot_beta_ga_accounts_chart.pptx"


def set_slide_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG


def add_textbox(slide, left, top, width, height, text, size=12, bold=False,
                color=COLOR_TEXT, align=PP_ALIGN.LEFT, italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.italic = italic
    p.font.color.rgb = color
    p.alignment = align
    return box


def _style_cell(cell, text, header=False, bold=False, color=None):
    cell.text = text
    p = cell.text_frame.paragraphs[0]
    p.font.name = FONT
    p.font.size = Pt(11 if header else 10)
    p.font.bold = header or bold
    if header:
        p.font.color.rgb = COLOR_WHITE
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_HEADER
    else:
        p.font.color.rgb = color or COLOR_TEXT


def build_chart_slide(slide):
    set_slide_bg(slide)
    if os.path.exists(OUTPUT_PNG):
        slide.shapes.add_picture(OUTPUT_PNG, Inches(1.4), Inches(0.35), width=Inches(10.5))
    add_textbox(
        slide, Inches(0.6), Inches(6.95), Inches(12), Inches(0.35),
        "Slide 2 has editable data · Upload to Google Drive → Open with Google Slides",
        size=9, color=COLOR_MID, italic=True, align=PP_ALIGN.CENTER,
    )


def build_data_slide(slide):
    set_slide_bg(slide)
    add_textbox(
        slide, Inches(0.6), Inches(0.4), Inches(12), Inches(0.5),
        "Accounts by stage — Pilot / Beta / GA (Editable)",
        size=24, bold=True, align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide, Inches(0.6), Inches(0.95), Inches(12), Inches(0.35),
        "July GA is for all accounts — leave the GA column blank / unlabeled in the chart.",
        size=11, color=COLOR_MID, italic=True, align=PP_ALIGN.CENTER,
    )

    headers = ["Month", "Pilot", "Beta", "GA", "Pilot+Beta total"]
    rows = []
    for month, pilot, beta in zip(MONTHS, PILOT, BETA):
        ga = "all accounts (no count)" if month == "July" else "—"
        rows.append([month, str(pilot), str(beta) if beta else "0", ga, str(pilot + beta)])

    table_shape = slide.shapes.add_table(
        len(rows) + 1, len(headers), Inches(1.0), Inches(1.6), Inches(11.3), Inches(2.8),
    )
    table = table_shape.table
    for i, w in enumerate([Inches(1.8), Inches(1.8), Inches(1.8), Inches(3.5), Inches(2.4)]):
        table.columns[i].width = w

    for ci, h in enumerate(headers):
        _style_cell(table.cell(0, ci), h, header=True)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            color = COLOR_GA if ci == 3 and "all" in val else None
            _style_cell(table.cell(ri, ci), val, bold=(ci == 0), color=color)

    add_textbox(
        slide, Inches(1.0), Inches(4.8), Inches(11), Inches(1.2),
        "How to mirror this in Google Sheets:\n"
        "1. Keep Pilot and Beta series with their numbers (and data labels).\n"
        "2. Add a third series named \"GA (all accounts)\" with blank/0 in May–June "
        "and a small placeholder in July only for the orange stack color.\n"
        "3. Turn OFF data labels for the GA series so no number appears.\n"
        "4. Optionally rename the July category to \"July (GA)\" or add a chart annotation.",
        size=11, color=COLOR_TEXT,
    )


def build_deck(output_path=OUTPUT_PPTX):
    if not os.path.exists(OUTPUT_PNG):
        create_chart()

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    s_chart = prs.slides.add_slide(blank)
    s_data = prs.slides.add_slide(blank)
    build_chart_slide(s_chart)
    build_data_slide(s_data)

    prs.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    build_deck()
